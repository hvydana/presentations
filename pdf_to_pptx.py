#!/usr/bin/env python3
"""Convert PDF to editable PPTX by extracting text blocks and images with positions."""
import sys, os, io
import fitz  # PyMuPDF
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

PDF_IN = "Physical_AI_SDK_v1.pdf"
PPTX_OUT = "Physical_AI_SDK_v1_edittable.pptx"

# 1 pt = 12700 EMU
PT_TO_EMU = 12700

doc = fitz.open(PDF_IN)
prs = Presentation()

# Set slide size to match the first page (points)
page0 = doc[0]
slide_w_pt = page0.rect.width
slide_h_pt = page0.rect.height
prs.slide_width = Emu(int(slide_w_pt * PT_TO_EMU))
prs.slide_height = Emu(int(slide_h_pt * PT_TO_EMU))

blank_layout = prs.slide_layouts[6]  # blank

img_dir = ".pdf_imgs_tmp"
os.makedirs(img_dir, exist_ok=True)

for pno, page in enumerate(doc):
    slide = prs.slides.add_slide(blank_layout)

    # --- Extract and place images ---
    img_list = page.get_images(full=True)
    placed_xrefs = set()
    for img_info in img_list:
        xref = img_info[0]
        if xref in placed_xrefs:
            continue
        placed_xrefs.add(xref)
        try:
            rects = page.get_image_rects(xref)
        except Exception:
            rects = []
        if not rects:
            continue
        try:
            pix = fitz.Pixmap(doc, xref)
            if pix.n - pix.alpha >= 4:  # CMYK
                pix = fitz.Pixmap(fitz.csRGB, pix)
            img_path = os.path.join(img_dir, f"p{pno}_x{xref}.png")
            pix.save(img_path)
            pix = None
        except Exception as e:
            print(f"img extract fail p{pno} xref{xref}: {e}")
            continue
        for r in rects:
            try:
                slide.shapes.add_picture(
                    img_path,
                    Emu(int(r.x0 * PT_TO_EMU)),
                    Emu(int(r.y0 * PT_TO_EMU)),
                    width=Emu(int((r.x1 - r.x0) * PT_TO_EMU)),
                    height=Emu(int((r.y1 - r.y0) * PT_TO_EMU)),
                )
            except Exception as e:
                print(f"img place fail: {e}")

    # --- Extract text blocks (editable) ---
    text_dict = page.get_text("dict")
    for block in text_dict.get("blocks", []):
        if block.get("type", 0) != 0:
            continue  # skip image blocks (we handled images above)
        bbox = block["bbox"]
        x0, y0, x1, y1 = bbox
        w = max(x1 - x0, 10)
        h = max(y1 - y0, 8)
        tb = slide.shapes.add_textbox(
            Emu(int(x0 * PT_TO_EMU)),
            Emu(int(y0 * PT_TO_EMU)),
            Emu(int(w * PT_TO_EMU)),
            Emu(int(h * PT_TO_EMU)),
        )
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Emu(0)
        tf.margin_top = tf.margin_bottom = Emu(0)
        first_para = True
        for line in block.get("lines", []):
            if first_para:
                p = tf.paragraphs[0]
                first_para = False
            else:
                p = tf.add_paragraph()
            for span in line.get("spans", []):
                run = p.add_run()
                run.text = span.get("text", "")
                sz = span.get("size", 11)
                run.font.size = Pt(sz)
                fname = span.get("font", "")
                if fname:
                    # strip subset prefix like "ABCDEF+"
                    if "+" in fname:
                        fname = fname.split("+", 1)[1]
                    run.font.name = fname
                flags = span.get("flags", 0)
                # bit 4 = bold, bit 1 = italic (PyMuPDF flags)
                if flags & 16:
                    run.font.bold = True
                if flags & 2:
                    run.font.italic = True
                color_int = span.get("color", 0)
                try:
                    r = (color_int >> 16) & 0xFF
                    g = (color_int >> 8) & 0xFF
                    b = color_int & 0xFF
                    run.font.color.rgb = RGBColor(r, g, b)
                except Exception:
                    pass

    print(f"page {pno+1}/{len(doc)} done")

prs.save(PPTX_OUT)
print(f"Wrote {PPTX_OUT}")